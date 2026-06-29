"""255172008_sunum.pptx — editlenebilir PowerPoint sunumu (Forest & Moss teması).

Beamer PDF ile aynı içerik; native pptx (PowerPoint'te düzenlenebilir).
Grafikler teslim/sunum/grafikler/ altından gömülür.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

TESLIM = Path(__file__).resolve().parent
GRAF = TESLIM / "sunum" / "grafikler"
OUT = TESLIM / "sunum" / "255172008_sunum.pptx"

FOREST = RGBColor(0x2C, 0x5F, 0x2D)
MOSS = RGBColor(0x97, 0xBC, 0x62)
CREAM = RGBColor(0xF5, 0xF5, 0xF5)
DARK = RGBColor(0x21, 0x21, 0x21)
GREY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HDR = "Cambria"
BODY = "Calibri"

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def _set(run, size, color, bold=False, italic=False, font=BODY):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    return tb, tf


def title(slide, text, color=FOREST, top=Inches(0.45)):
    tb, tf = textbox(slide, Inches(0.6), top, Inches(12.1), Inches(0.95))
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    _set(r, 30, color, bold=True, font=HDR)
    return tb


def bullets(slide, items, left=Inches(0.7), top=Inches(1.7), width=Inches(12.0), height=Inches(5.2), size=18, gap=10):
    tb, tf = textbox(slide, left, top, width, height)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        if isinstance(item, tuple):  # (bold_head, rest)
            head, rest = item
            r = p.add_run(); r.text = "▸ " + head; _set(r, size, FOREST, bold=True)
            if rest:
                r2 = p.add_run(); r2.text = " " + rest; _set(r2, size, DARK)
        else:
            r = p.add_run(); r.text = "▸ " + item; _set(r, size, DARK)
    return tb


def content_slide(t, items, **kw):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    title(s, t)
    bullets(s, items, **kw)
    return s


def image_slide(t, img, caption):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    title(s, t)
    # Görseli oranını koruyarak alana sığdır.
    from PIL import Image
    iw, ih = Image.open(img).size
    avail_w, avail_h = Inches(11.8), Inches(4.9)
    scale = min(avail_w / iw, avail_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    left = int((SW - w) / 2)
    s.shapes.add_picture(str(img), left, Inches(1.55), width=w, height=h)
    tb, tf = textbox(s, Inches(0.7), Inches(6.75), Inches(12.0), Inches(0.55))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = caption; _set(r, 12, GREY, italic=True)
    return s


# === 1. Başlık (koyu) ===
s = prs.slides.add_slide(BLANK); bg(s, FOREST)
_, tf = textbox(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.4), anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Süt Sığırcılığında İnek Bazında Keep/Cull Kararı için"
_set(r, 34, WHITE, bold=True, font=HDR)
p2 = tf.add_paragraph(); r = p2.add_run(); r.text = "Pekiştirmeli Öğrenme"
_set(r, 34, WHITE, bold=True, font=HDR)
p3 = tf.add_paragraph(); p3.space_before = Pt(14)
r = p3.add_run(); r.text = "Empirik Stokastik MDP  ·  Tabular Q-Learning + SB3 DQN"
_set(r, 18, MOSS, italic=True)
_, tf2 = textbox(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.7))
r = tf2.paragraphs[0].add_run(); r.text = "255172008 — Artun KARA"
_set(r, 20, WHITE, bold=True)

# === 2. Problem + Neden RL ===
content_slide("Problem ve Neden Pekiştirmeli Öğrenme?", [
    ("Problem:", "Her ay her inek için keep (elde tut) / cull (sürüden çıkar); amaç iskontolu uzun vadeli iktisadi getiriyi maksimize etmek."),
    ("Ardışıklık:", "keep ineği bir sonraki aylık duruma taşır; karar geleceği etkiler."),
    ("Gecikmeli ödül:", "Tutmanın değeri gelecekteki süt − birikmiş risk (γ = 0.97)."),
    ("Stokastik + büyük uzay:", "Geçiş dinamiği gerçek kayıtlardan empirik; 162 durum."),
    ("Neden A*/DP değil:", "Ortam stokastik (kod satırıyla ispatlı) → en kısa yol uygulanamaz; risk-eşiği heuristiği RL'in gerisinde."),
])

# === 3. MDP Formülasyonu ===
s = content_slide("MDP Formülasyonu", [
    ("Durum:", "laktasyon_evresi | parite | dim_band | milk_level | milk_trend | health_risk | risk_band (162 kova)."),
    ("Aksiyon:", "keep (0, terminal değil), cull (1, terminal)."),
    ("Ödül (kod ↔ formül birebir):", ""),
], height=Inches(2.6))
# Ödül formülü kutusu
_, tf = textbox(s, Inches(0.9), Inches(4.0), Inches(11.5), Inches(2.6))
for txt, code in [
    ("R_keep = m·d·p − c_b·d − C_sağlık − C_üreme − C_beklenen-cull", "kod/env/rl_mdp_utils.py:120"),
    ("R_cull = v_c − c_r − c_disposal   (c_disposal = 0)", "kod/env/rl_mdp_utils.py:76"),
]:
    p = tf.add_paragraph() if tf.paragraphs[0].runs else tf.paragraphs[0]
    p.space_after = Pt(8)
    r = p.add_run(); r.text = txt; _set(r, 18, DARK, bold=True, font="Consolas")
    r2 = p.add_run(); r2.text = "    (" + code + ")"; _set(r2, 13, GREY, italic=True)
p = tf.add_paragraph(); p.space_before = Pt(6)
r = p.add_run(); r.text = "Parametreler reward_baseline_us_2025.csv (süt fiyatı & bakım USDA 2025; diğerleri prototip)."
_set(r, 13, GREY, italic=True)

# === 4. Ortam Stokastik ===
s = content_slide("Ortam Stokastik (kanıt)", [
    ("162 durumun 120'sinde (%74)", "keep altında birden çok olası s′ var."),
    ("Örnek:", "bir durum 3 farklı s′'ye her biri p = 0.333 ile geçer."),
    ("cull deterministik terminal.", "Aynı seed → aynı yörünge (yalnızca tekrarlanabilirlik)."),
], top=Inches(3.7), height=Inches(3.2))
_, tf = textbox(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(1.7))
p = tf.paragraphs[0]
r = p.add_run(); r.text = "keep sonrası bir sonraki durum empirik dağılımdan örneklenir:"
_set(r, 17, DARK)
p2 = tf.add_paragraph(); p2.space_before = Pt(8)
r = p2.add_run(); r.text = "selected = outcomes[int(rng.choice(len(outcomes), p=probabilities))]"
_set(r, 15, FOREST, bold=True, font="Consolas")
p3 = tf.add_paragraph()
r = p3.add_run(); r.text = "kod/env/rl_experiment_utils.py:161"
_set(r, 13, GREY, italic=True)

# === 5. Yöntem + Kurulum (iki sütun) ===
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "Yöntem ve Deney Kurulumu")
_, tfl = textbox(s, Inches(0.7), Inches(1.7), Inches(6.0), Inches(5.0))
r = tfl.paragraphs[0].add_run(); r.text = "Ajanlar / referans"; _set(r, 20, FOREST, bold=True, font=HDR)
for it in ["Tabular Q-Learning (off-policy)", "SB3 DQN (off-policy, replay buffer)",
           "MDP Value Iteration (optimum ref.)", "Baseline: random / keep_all / cull_all / risk_threshold"]:
    p = tfl.add_paragraph(); p.space_after = Pt(8); r = p.add_run(); r.text = "▸ " + it; _set(r, 16, DARK)
_, tfr = textbox(s, Inches(7.0), Inches(1.7), Inches(5.6), Inches(5.0))
r = tfr.paragraphs[0].add_run(); r.text = "Kurulum"; _set(r, 20, FOREST, bold=True, font=HDR)
for it in ["Seedler: 42, 123, 2024, 7, 99 (≥5)", "full: DQN 100k, Q-learning 4000 ep.",
           "np.random.default_rng (yasak kalıp yok)", "Eval: deterministic/greedy, 500 ep/seed"]:
    p = tfr.add_paragraph(); p.space_after = Pt(8); r = p.add_run(); r.text = "▸ " + it; _set(r, 16, DARK)

# === 6. Sonuçlar tablosu ===
s = prs.slides.add_slide(BLANK); bg(s, WHITE); title(s, "Sonuçlar (seed ort. ± std, full)")
rows = [
    ("Politika", "Tür", "Getiri (USD)", "Std", "Cull"),
    ("qlearning", "ajan", "−204", "30", "0.72"),
    ("mdp_value_iteration", "referans", "−225", "42", "0.45"),
    ("dqn", "ajan", "−270", "20", "0.82"),
    ("risk_threshold", "baseline", "−385", "53", "0.85"),
    ("keep_all", "baseline", "−428", "34", "0.00"),
    ("random_policy", "baseline", "−964", "56", "0.75"),
    ("cull_all", "baseline", "−1300", "0", "1.00"),
]
tbl = s.shapes.add_table(len(rows), 5, Inches(1.4), Inches(1.6), Inches(10.5), Inches(4.3)).table
tbl.columns[0].width = Inches(3.6)
for c in range(1, 5):
    tbl.columns[c].width = Inches(1.725)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.text = val
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
        run = para.runs[0]
        if ri == 0:
            _set(run, 14, WHITE, bold=True); cell.fill.solid(); cell.fill.fore_color.rgb = FOREST
        else:
            is_agent = row[1] == "ajan"
            _set(run, 13, DARK, bold=is_agent)
            cell.fill.solid()
            cell.fill.fore_color.rgb = MOSS if is_agent else (CREAM if ri % 2 else WHITE)
_, tf = textbox(s, Inches(0.7), Inches(6.2), Inches(12.0), Inches(0.9))
r = tf.paragraphs[0].add_run()
r.text = "Ajanlar tüm baseline'ları geçti; Q-learning model-tabanlı optimuma istatistiksel olarak eşitlendi (fark anlamsız, eşli t-test p≈0.37). Dağılım sağa çarpık: medyan −556 vs ort. −204."
_set(r, 14, DARK, italic=True)

# === 7-11. Grafikler ===
image_slide("Grafik 1 — Öğrenme Eğrisi", GRAF / "01_ogrenme_egrisi.png",
            "y: episode getirisi (anlık ödül değil); n=5 seed, ortalama ± std + hareketli ortalama.")
image_slide("Grafik 2 — Test/Eval Eğrisi", GRAF / "02_eval_egrisi.png",
            "Deterministic/greedy; yeşil kesik çizgi MDP-optimum referansı (tavan). Ajanlar tavana ulaşıyor.")
image_slide("Grafik 3 — Loss + Policy Heatmap", GRAF / "03_loss_ve_policy_heatmap.png",
            "DQN TD-loss yakınsıyor; heatmap: yüksek risk × ileri DIM → cull oranı artıyor (iktisadi olarak anlamlı).")
image_slide("Grafik 4 — Hiperparametre Duyarlılığı", GRAF / "04_hiperparametre_duyarlilik.png",
            "Sweet-spot: orta keşif × orta öğrenme oranı (ε₀=0.35, α=0.18) = seçilen varsayılanlar.")
image_slide("Grafik 5 — Baseline Karşılaştırma", GRAF / "05_baseline_karsilastirma.png",
            "Eğitilmiş ajanlar (mavi) tüm naif/heuristik baseline'ları (turuncu) geçiyor; MDP optimum yeşil referans.")

# === 12. Savunma özeti ===
content_slide("Teslim-3 Savunma — Özet (kod satırlı)", [
    ("Stokastik:", "keep için empirik p(s′|s,a) — rl_experiment_utils.py:161; %74 durumda >1 olası s′."),
    ("Off-policy:", "Q-learning davranış ε-greedy (rl_core.py:309) vs hedef max Q (:321); DQN replay buffer (:474). Online ≠ offline."),
    ("Markov:", "state_bucket → next varsayımı; trend/risk geçmişi özetler → yaklaşık Markov (bucketing kaybı)."),
    ("POMDP:", "cull_risk = denetimli tahmin (gerçek değil); feature engineering ile gözlemi Markov'a yaklaştırma."),
], size=17)

# === 13. Sonuç (koyu) ===
s = prs.slides.add_slide(BLANK); bg(s, FOREST)
title(s, "Sonuç", color=WHITE)
_, tf = textbox(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(4.0))
for it in ["Empirik, stokastik, gecikmeli-ödüllü keep/cull ortamı → RL gerekçeli.",
           "Q-Learning ve DQN, ≥5 seed üzerinde tüm baseline'ları geçti; model-tabanlı optimuma istatistiksel olarak eşitlendi (fark anlamsız).",
           "Tüm sonuçlar tek komutla (run_all.sh) yeniden üretilebilir; ham loglardan doğrulanabilir."]:
    p = tf.add_paragraph() if tf.paragraphs[0].runs else tf.paragraphs[0]
    p.space_after = Pt(14); r = p.add_run(); r.text = "▸ " + it; _set(r, 19, WHITE)
_, tf2 = textbox(s, Inches(0.9), Inches(6.0), Inches(11.5), Inches(0.9))
p = tf2.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Teşekkürler."; _set(r, 26, MOSS, bold=True, font=HDR)

n_slides = len(prs.slides._sldIdLst)
prs.save(OUT)
print(f"Yazıldı: {OUT.relative_to(TESLIM.parent)}  ({n_slides} slayt)")
