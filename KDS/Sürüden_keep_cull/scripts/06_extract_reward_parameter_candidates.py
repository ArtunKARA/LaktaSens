from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ERS_DIR = PROJECT_ROOT / "data" / "external" / "usda_ers_milk_cost"
NASS_DIR = PROJECT_ROOT / "data" / "external" / "usda_nass_milk_price"
OUTPUT_CSV = PROJECT_ROOT / "data" / "processed" / "reward_parameter_source_candidates.csv"
OUTPUT_MD = PROJECT_ROOT / "reports" / "economic_reward_parameter_review.md"

FIELDNAMES = [
    "source_code",
    "source_file",
    "candidate_param",
    "year",
    "period",
    "geography",
    "size",
    "category",
    "item",
    "units",
    "value_raw",
    "value_status",
    "notes",
]

ERS_CANDIDATE_ITEMS = {
    "Milk sold": "sut_fiyati_reference_per_cwt",
    "Total, costs listed": "total_costs_listed_per_cwt",
    "Total, operating costs": "operating_costs_per_cwt",
    "Total, feed costs": "feed_costs_per_cwt",
    "Output per cow": "output_per_cow_reference",
    "Milk cows": "milk_cows_reference",
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_ers_candidates() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in sorted(ERS_DIR.glob("*.csv")):
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                with path.open("r", encoding=encoding, newline="") as file:
                    reader = csv.DictReader(file)
                    for row in reader:
                        item = clean_text(row.get("Item"))
                        candidate_param = ERS_CANDIDATE_ITEMS.get(item)
                        if not candidate_param:
                            continue
                        rows.append(
                            {
                                "source_code": "usda_ers_milk_cost",
                                "source_file": str(path.relative_to(PROJECT_ROOT)),
                                "candidate_param": candidate_param,
                                "year": clean_text(row.get("Year")),
                                "period": "YEAR",
                                "geography": clean_text(row.get("Region")),
                                "size": clean_text(row.get("Size")),
                                "category": clean_text(row.get("Category")),
                                "item": item,
                                "units": clean_text(row.get("Units")),
                                "value_raw": clean_text(row.get("Value")),
                                "value_status": "ok_from_csv",
                                "notes": "Candidate only. Convert units and choose accounting convention before inserting into feat_reward_parametre.",
                            }
                        )
                break
            except UnicodeDecodeError:
                rows = [
                    row
                    for row in rows
                    if row["source_file"] != str(path.relative_to(PROJECT_ROOT))
                ]
                continue
        else:
            raise UnicodeDecodeError(
                "csv",
                b"",
                0,
                1,
                f"Could not decode {path} with supported encodings.",
            )
    return rows


def value_status(value: Any) -> tuple[str, str]:
    if isinstance(value, datetime | date):
        return (
            value.isoformat(),
            "invalid_excel_date_coercion",
        )
    if value is None or clean_text(value) == "":
        return "", "missing"
    text_value = clean_text(value).replace(",", "")
    try:
        float(text_value)
        return text_value, "ok_from_xlsx"
    except ValueError:
        return clean_text(value), "manual_review_non_numeric"


def read_nass_candidates() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in sorted(NASS_DIR.glob("*.xlsx")):
        workbook = load_workbook(path, read_only=True, data_only=True)
        for worksheet in workbook.worksheets:
            records = worksheet.iter_rows(values_only=True)
            try:
                headers = [clean_text(value) for value in next(records)]
            except StopIteration:
                continue
            for record in records:
                row = dict(zip(headers, record))
                data_item = clean_text(row.get("Data Item"))
                if "PRICE RECEIVED" not in data_item.upper():
                    continue
                raw_value, status = value_status(row.get("Value"))
                rows.append(
                    {
                        "source_code": "usda_nass_milk_price",
                        "source_file": str(path.relative_to(PROJECT_ROOT)),
                        "candidate_param": "sut_fiyati_per_cwt",
                        "year": clean_text(row.get("Year")),
                        "period": clean_text(row.get("Period")),
                        "geography": clean_text(row.get("State")) or clean_text(row.get("Geo Level")),
                        "size": "",
                        "category": "PRICE RECEIVED",
                        "item": data_item,
                        "units": "$ / CWT",
                        "value_raw": raw_value,
                        "value_status": status,
                        "notes": "Do not insert if value_status is invalid_excel_date_coercion; re-download from NASS as CSV or fix source export.",
                    }
                )
    return rows


def write_csv(rows: list[dict[str, str]]) -> None:
    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_CSV.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)


def write_markdown(rows: list[dict[str, str]]) -> None:
    OUTPUT_MD.parent.mkdir(parents=True, exist_ok=True)
    by_source: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        by_source.setdefault(row["source_code"], []).append(row)

    invalid_nass = [
        row
        for row in rows
        if row["source_code"] == "usda_nass_milk_price"
        and row["value_status"] == "invalid_excel_date_coercion"
    ]
    nass_interpretation = (
        "`usda_nass_milk_price` should feed `sut_fiyati`; the current file has clean numeric/string price values."
        if not invalid_nass
        else "`usda_nass_milk_price` should feed `sut_fiyati`, but the current XLSX has date-coerced values in some rows. Re-download as CSV or fix the source export before using it."
    )

    lines = [
        "# Economic Reward Parameter Review",
        "",
        "This report lists source-level candidate economic values. Candidate rows are not inserted directly into `feat_reward_parametre`; the reviewed prototype baseline is generated separately by `scripts/11_build_reward_baseline.py` after unit conversion and scenario assumptions are documented.",
        "",
        "## Summary",
        "",
        f"- Total candidate rows: `{len(rows)}`",
        f"- ERS candidate rows: `{len(by_source.get('usda_ers_milk_cost', []))}`",
        f"- NASS candidate rows: `{len(by_source.get('usda_nass_milk_price', []))}`",
        f"- NASS rows with Excel date coercion: `{len(invalid_nass)}`",
        "",
        "## Interpretation",
        "",
        f"- {nass_interpretation}",
        "- `usda_ers_milk_cost` contains annual cost/revenue estimates by state and by operation size. The `Value` field is usually in dollars per hundredweight sold for cost rows, so conversion is required before deriving cow-day or cow-month costs.",
        "- `MilkState2021base.csv` is the better first baseline file for a single national/state scenario. `MilkSize2021base.csv` is useful later for size-specific sensitivity analysis.",
        "",
        "## Candidate Parameters",
        "",
        "| Source | Candidate | Rows | Notes |",
        "|---|---:|---:|---|",
    ]

    counts: dict[tuple[str, str], int] = {}
    for row in rows:
        key = (row["source_code"], row["candidate_param"])
        counts[key] = counts.get(key, 0) + 1
    for (source_code, candidate), count in sorted(counts.items()):
        note = "manual unit conversion required"
        if source_code == "usda_nass_milk_price":
            note = "clean current file; still requires scenario/date convention review"
            if invalid_nass:
                note = "current XLSX has date-coerced rows; fix before use"
        lines.append(f"| `{source_code}` | `{candidate}` | `{count}` | {note} |")

    OUTPUT_MD.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> int:
    rows = read_ers_candidates() + read_nass_candidates()
    write_csv(rows)
    write_markdown(rows)
    print(f"wrote {OUTPUT_CSV.relative_to(PROJECT_ROOT)} ({len(rows)} rows)")
    print(f"wrote {OUTPUT_MD.relative_to(PROJECT_ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
