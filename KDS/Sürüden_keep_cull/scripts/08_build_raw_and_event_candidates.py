from __future__ import annotations

import csv
import json
import zipfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXTERNAL_DIR = PROJECT_ROOT / "data" / "external"
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"
RAW_DIR = PROCESSED_DIR / "raw_candidates"
EVENT_DIR = PROCESSED_DIR / "event_candidates"
EVENT_NORMALIZED_PATH = PROCESSED_DIR / "event_normalized_candidates.csv"
REJECTED_PATH = PROCESSED_DIR / "rejected_rows.csv"
MAP_PATH = PROCESSED_DIR / "map_external_inek_candidates.csv"

EVENT_TYPES = {
    "calving",
    "insemination",
    "pregnancy_diagnosis_positive",
    "pregnancy_loss",
    "dryoff",
    "cull",
    "milk_measurement",
    "health_lameness",
    "group_change",
    "entry",
    "birth",
    "other",
}

RAW_FIELDS = {
    "raw_repro_loss": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "fparto",
        "nparto",
        "nlact",
        "nrosf",
        "fechasf",
        "fechaperd",
        "npre",
        "perdida",
        "tperd",
        "codigo_tperd",
        "estsf",
        "dgp",
        "delp",
        "sana_ribeiro",
        "cl100",
        "cl150",
        "clsf",
        "cltacto",
        "raw_payload",
    ],
    "raw_milk": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "external_machine_id",
        "sagim_tarihi",
        "sut_miktar_kg",
        "sagim_no",
        "conductivity",
        "yag_orani",
        "protein_orani",
        "scc",
        "raw_payload",
    ],
    "raw_culling": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "cikarilma_tarihi",
        "cikarilma_sebebi",
        "planned_cull",
        "satis_degeri",
        "raw_payload",
    ],
    "raw_health": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "olay_tarihi",
        "olay_tipi",
        "alt_tip",
        "siddet",
        "tedavi_verildi",
        "aciklama",
        "raw_payload",
    ],
}

EVENT_FIELDS = {
    "event_dogurma": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "dogurma_tarihi",
        "parite",
        "laktasyon_no",
        "source_table",
        "source_candidate_id",
        "payload",
    ],
    "event_dollenme": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "dollenme_tarihi",
        "dollenme_sayisi",
        "tuttu_mu",
        "servis_durumu",
        "source_table",
        "source_candidate_id",
        "payload",
    ],
    "event_kuru": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "kuru_baslangic_tarihi",
        "source_table",
        "source_candidate_id",
        "payload",
    ],
    "event_sut_olcum": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "sagim_tarihi",
        "sut_miktar_kg",
        "sagim_no",
        "conductivity",
        "yag_orani",
        "protein_orani",
        "scc",
        "source_table",
        "source_candidate_id",
        "payload",
    ],
    "event_suruden_cikarma": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "cikarilma_tarihi",
        "cikarilma_sebebi",
        "planned_cull",
        "satis_degeri",
        "source_table",
        "source_candidate_id",
        "payload",
    ],
    "event_health": [
        "candidate_id",
        "source_code",
        "file_name",
        "sheet_name",
        "source_row_number",
        "external_animal_id",
        "canonical_animal_id",
        "olay_tarihi",
        "olay_tipi",
        "alt_tip",
        "siddet",
        "tedavi_verildi",
        "aciklama",
        "source_table",
        "source_candidate_id",
        "payload",
    ],
}

NORMALIZED_FIELDS = [
    "candidate_id",
    "source_code",
    "file_name",
    "sheet_name",
    "source_row_number",
    "external_animal_id",
    "canonical_animal_id",
    "event_date",
    "event_type",
    "source_table",
    "source_candidate_id",
    "numeric_value",
    "text_value",
    "payload",
]

REJECT_FIELDS = [
    "source_code",
    "file_name",
    "sheet_name",
    "source_row_number",
    "target_table",
    "reason",
    "row_payload",
]

MAP_FIELDS = ["source_code", "external_animal_id", "canonical_animal_id", "notlar"]


def to_jsonable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    return "" if text.lower() in {"none", "nan", "nat"} else text


def payload_json(row: dict[str, Any]) -> str:
    return json.dumps({key: to_jsonable(value) for key, value in row.items()}, ensure_ascii=False, default=str)


def parse_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        return ""


def parse_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def number_text(value: Any, decimals: int | None = None) -> str:
    parsed = parse_float(value)
    if parsed is None:
        return ""
    if decimals is None:
        return str(parsed)
    return f"{parsed:.{decimals}f}"


def truthy(value: Any) -> bool:
    text = clean(value).lower()
    return text in {"1", "true", "t", "yes", "y", "e", "evet", "positive", "+", "si", "sí"}


def canonical_animal_id(source_code: str, external_animal_id: Any) -> str:
    return f"{source_code}:{clean(external_animal_id)}"


def unique_headers(raw_headers: Iterable[Any]) -> list[str]:
    seen: dict[str, int] = {}
    output: list[str] = []
    for index, value in enumerate(raw_headers, start=1):
        base = clean(value) or f"unnamed_{index}"
        seen[base] = seen.get(base, 0) + 1
        output.append(base if seen[base] == 1 else f"{base}__{seen[base]}")
    return output


def iter_xlsx_rows_from_zip(zip_path: Path, member_name: str, sheet_name: str) -> Iterable[tuple[int, dict[str, Any]]]:
    with zipfile.ZipFile(zip_path) as zip_file:
        with zip_file.open(member_name) as raw:
            workbook = load_workbook(BytesIO(raw.read()), read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = unique_headers(next(iterator))
    except StopIteration:
        return
    for row_number, values in enumerate(iterator, start=2):
        if values is None or not any(clean(value) for value in values):
            continue
        row = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        yield row_number, row


class CandidateOutputs:
    def __init__(self) -> None:
        RAW_DIR.mkdir(parents=True, exist_ok=True)
        EVENT_DIR.mkdir(parents=True, exist_ok=True)
        PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
        self.files: list[Any] = []
        self.writers: dict[str, csv.DictWriter] = {}
        self.counts: dict[str, int] = {}
        for name, fields in RAW_FIELDS.items():
            self._open_writer(name, RAW_DIR / f"{name}.csv", fields)
        for name, fields in EVENT_FIELDS.items():
            self._open_writer(name, EVENT_DIR / f"{name}.csv", fields)
        self._open_writer("event_normalized", EVENT_NORMALIZED_PATH, NORMALIZED_FIELDS)
        self._open_writer("rejected_rows", REJECTED_PATH, REJECT_FIELDS)

    def _open_writer(self, name: str, path: Path, fieldnames: list[str]) -> None:
        file = path.open("w", encoding="utf-8", newline="")
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        self.files.append(file)
        self.writers[name] = writer
        self.counts[name] = 0

    def close(self) -> None:
        for file in self.files:
            file.close()

    def write(self, name: str, row: dict[str, Any]) -> str:
        writer = self.writers[name]
        self.counts[name] += 1
        candidate_id = row.get("candidate_id") or f"{name}:{self.counts[name]}"
        row = {**row, "candidate_id": candidate_id}
        completed = {field: clean(row.get(field, "")) for field in writer.fieldnames or []}
        writer.writerow(completed)
        return str(candidate_id)

    def reject(
        self,
        source_code: str,
        file_name: str,
        sheet_name: str,
        source_row_number: int | str,
        target_table: str,
        reason: str,
        row: dict[str, Any],
    ) -> None:
        self.write(
            "rejected_rows",
            {
                "source_code": source_code,
                "file_name": file_name,
                "sheet_name": sheet_name,
                "source_row_number": source_row_number,
                "target_table": target_table,
                "reason": reason,
                "row_payload": payload_json(row),
            },
        )


def add_map(map_records: set[tuple[str, str, str]], source_code: str, external_id: Any) -> str:
    external = clean(external_id)
    if not external:
        return ""
    canonical = canonical_animal_id(source_code, external)
    map_records.add((source_code, external, canonical))
    return canonical


def add_normalized(
    outputs: CandidateOutputs,
    source_code: str,
    file_name: str,
    sheet_name: str,
    source_row_number: int | str,
    external_animal_id: str,
    event_date: str,
    event_type: str,
    source_table: str,
    source_candidate_id: str,
    row: dict[str, Any],
    numeric_value: Any = "",
    text_value: Any = "",
) -> None:
    normalized_type = event_type if event_type in EVENT_TYPES else "other"
    outputs.write(
        "event_normalized",
        {
            "source_code": source_code,
            "file_name": file_name,
            "sheet_name": sheet_name,
            "source_row_number": source_row_number,
            "external_animal_id": external_animal_id,
            "canonical_animal_id": canonical_animal_id(source_code, external_animal_id),
            "event_date": event_date,
            "event_type": normalized_type,
            "source_table": source_table,
            "source_candidate_id": source_candidate_id,
            "numeric_value": numeric_value,
            "text_value": text_value,
            "payload": payload_json(row),
        },
    )


def process_embryo(outputs: CandidateOutputs, map_records: set[tuple[str, str, str]]) -> None:
    source_code = "embryo_losses_mendeley"
    zip_path = EXTERNAL_DIR / source_code / "EmbryoLosses.zip"
    member = "EmbryoLosses/CorvaSG et al.xlsx"
    sheet = "Hoja1"
    file_name = f"{zip_path.relative_to(PROJECT_ROOT)}::{member}"
    if not zip_path.exists():
        outputs.reject(source_code, str(zip_path), sheet, "", "raw_repro_loss", "source file missing", {})
        return

    raw_columns = [
        "fparto",
        "nparto",
        "nlact",
        "nrosf",
        "fechasf",
        "fechaperd",
        "npre",
        "perdida",
        "tperd",
        "codigo_tperd",
        "estsf",
        "dgp",
        "delp",
        "sana_ribeiro",
        "cl100",
        "cl150",
        "clsf",
        "cltacto",
    ]
    for row_number, row in iter_xlsx_rows_from_zip(zip_path, member, sheet):
        external_id = clean(row.get("id"))
        if not external_id:
            outputs.reject(source_code, file_name, sheet, row_number, "raw_repro_loss", "missing id", row)
            continue
        canonical = add_map(map_records, source_code, external_id)
        raw_row = {
            "source_code": source_code,
            "file_name": file_name,
            "sheet_name": sheet,
            "source_row_number": row_number,
            "external_animal_id": external_id,
            "canonical_animal_id": canonical,
            "raw_payload": payload_json(row),
        }
        for column in raw_columns:
            raw_row[column] = parse_date(row.get(column)) if column in {"fparto", "fechasf", "fechaperd"} else clean(row.get(column))
        raw_candidate_id = outputs.write("raw_repro_loss", raw_row)

        fparto = parse_date(row.get("fparto"))
        if fparto:
            event_id = outputs.write(
                "event_dogurma",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "dogurma_tarihi": fparto,
                    "parite": row.get("nparto"),
                    "laktasyon_no": row.get("nlact"),
                    "source_table": "raw_repro_loss",
                    "source_candidate_id": raw_candidate_id,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, fparto, "calving", "event_dogurma", event_id, row)

        fechasf = parse_date(row.get("fechasf"))
        if fechasf:
            event_id = outputs.write(
                "event_dollenme",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "dollenme_tarihi": fechasf,
                    "dollenme_sayisi": row.get("nrosf"),
                    "servis_durumu": row.get("estsf"),
                    "source_table": "raw_repro_loss",
                    "source_candidate_id": raw_candidate_id,
                    "payload": payload_json(row),
                },
            )
            add_normalized(
                outputs,
                source_code,
                file_name,
                sheet,
                row_number,
                external_id,
                fechasf,
                "insemination",
                "event_dollenme",
                event_id,
                row,
                text_value=row.get("estsf"),
            )

        loss_date = parse_date(row.get("fechaperd"))
        if truthy(row.get("perdida")) and loss_date:
            add_normalized(
                outputs,
                source_code,
                file_name,
                sheet,
                row_number,
                external_id,
                loss_date,
                "pregnancy_loss",
                "raw_repro_loss",
                raw_candidate_id,
                row,
                text_value=row.get("codigo_tperd"),
            )


def process_primiparous_daily(outputs: CandidateOutputs, map_records: set[tuple[str, str, str]]) -> None:
    source_code = "primiparous_social_env"
    zip_path = EXTERNAL_DIR / source_code / "primiparous social environment.zip"
    member = "j42n4x723g-2/data in Brief publish.xlsx"
    sheet = "Daily data"
    file_name = f"{zip_path.relative_to(PROJECT_ROOT)}::{member}"
    if not zip_path.exists():
        outputs.reject(source_code, str(zip_path), sheet, "", "raw_milk", "source file missing", {})
        return

    for row_number, row in iter_xlsx_rows_from_zip(zip_path, member, sheet):
        external_id = clean(row.get("CowID"))
        measurement_date = parse_date(row.get("Date"))
        yield_grams = parse_float(row.get("Yield(gr)"))
        if not external_id:
            outputs.reject(source_code, file_name, sheet, row_number, "raw_milk", "missing CowID", row)
            continue
        if not measurement_date:
            outputs.reject(source_code, file_name, sheet, row_number, "raw_milk", "missing Date", row)
            continue
        if yield_grams is None:
            outputs.reject(source_code, file_name, sheet, row_number, "raw_milk", "missing Yield(gr)", row)
            continue

        canonical = add_map(map_records, source_code, external_id)
        milk_kg = yield_grams / 1000
        raw_candidate_id = outputs.write(
            "raw_milk",
            {
                "source_code": source_code,
                "file_name": file_name,
                "sheet_name": sheet,
                "source_row_number": row_number,
                "external_animal_id": external_id,
                "canonical_animal_id": canonical,
                "sagim_tarihi": measurement_date,
                "sut_miktar_kg": f"{milk_kg:.6f}",
                "conductivity": number_text(row.get("Conductivity")),
                "raw_payload": payload_json(row),
            },
        )
        event_id = outputs.write(
            "event_sut_olcum",
            {
                "source_code": source_code,
                "file_name": file_name,
                "sheet_name": sheet,
                "source_row_number": row_number,
                "external_animal_id": external_id,
                "canonical_animal_id": canonical,
                "sagim_tarihi": measurement_date,
                "sut_miktar_kg": f"{milk_kg:.6f}",
                "conductivity": number_text(row.get("Conductivity")),
                "source_table": "raw_milk",
                "source_candidate_id": raw_candidate_id,
                "payload": payload_json(row),
            },
        )
        add_normalized(
            outputs,
            source_code,
            file_name,
            sheet,
            row_number,
            external_id,
            measurement_date,
            "milk_measurement",
            "event_sut_olcum",
            event_id,
            row,
            numeric_value=f"{milk_kg:.6f}",
        )


def process_primiparous_events(outputs: CandidateOutputs, map_records: set[tuple[str, str, str]]) -> None:
    source_code = "primiparous_social_env"
    zip_path = EXTERNAL_DIR / source_code / "primiparous social environment.zip"
    member = "j42n4x723g-2/data in Brief publish.xlsx"
    sheet = "events"
    file_name = f"{zip_path.relative_to(PROJECT_ROOT)}::{member}"
    if not zip_path.exists():
        outputs.reject(source_code, str(zip_path), sheet, "", "event_normalized", "source file missing", {})
        return

    for row_number, row in iter_xlsx_rows_from_zip(zip_path, member, sheet):
        external_id = clean(row.get("CowID"))
        event_date = parse_date(row.get("Date"))
        event_name = clean(row.get("Event"))
        if not external_id:
            outputs.reject(source_code, file_name, sheet, row_number, "event_normalized", "missing CowID", row)
            continue
        if not event_date:
            outputs.reject(source_code, file_name, sheet, row_number, "event_normalized", "missing Date", row)
            continue
        if not event_name:
            outputs.reject(source_code, file_name, sheet, row_number, "event_normalized", "missing Event", row)
            continue

        canonical = add_map(map_records, source_code, external_id)
        event_lower = event_name.lower()
        status = clean(row.get("Status"))
        group = clean(row.get("Group"))

        if event_lower == "calving":
            event_id = outputs.write(
                "event_dogurma",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "dogurma_tarihi": event_date,
                    "source_table": "source_events_sheet",
                    "source_candidate_id": row_number,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, event_date, "calving", "event_dogurma", event_id, row, text_value=status)
        elif event_lower == "insemination":
            event_id = outputs.write(
                "event_dollenme",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "dollenme_tarihi": event_date,
                    "servis_durumu": status or event_name,
                    "source_table": "source_events_sheet",
                    "source_candidate_id": row_number,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, event_date, "insemination", "event_dollenme", event_id, row, text_value=status)
        elif event_lower == "dry":
            event_id = outputs.write(
                "event_kuru",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "kuru_baslangic_tarihi": event_date,
                    "source_table": "source_events_sheet",
                    "source_candidate_id": row_number,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, event_date, "dryoff", "event_kuru", event_id, row, text_value=status)
        elif event_lower == "exit":
            raw_id = outputs.write(
                "raw_culling",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "cikarilma_tarihi": event_date,
                    "cikarilma_sebebi": status or event_name,
                    "raw_payload": payload_json(row),
                },
            )
            event_id = outputs.write(
                "event_suruden_cikarma",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "cikarilma_tarihi": event_date,
                    "cikarilma_sebebi": status or event_name,
                    "source_table": "raw_culling",
                    "source_candidate_id": raw_id,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, event_date, "cull", "event_suruden_cikarma", event_id, row, text_value=status)
        else:
            normalized_type = {
                "pd (+)": "pregnancy_diagnosis_positive",
                "change group": "group_change",
                "entry": "entry",
                "birth": "birth",
                "abortion": "pregnancy_loss",
            }.get(event_lower, "other")
            add_normalized(
                outputs,
                source_code,
                file_name,
                sheet,
                row_number,
                external_id,
                event_date,
                normalized_type,
                "source_events_sheet",
                row_number,
                row,
                text_value=status or group or event_name,
            )


def process_chiozza(outputs: CandidateOutputs, map_records: set[tuple[str, str, str]]) -> None:
    source_code = "chiozza_lameness"
    zip_path = EXTERNAL_DIR / source_code / "Chiozza.zip"
    member = "ChiozzaLogroñoJ_SupplementaryData.xlsx"
    sheet = "Hoja1"
    file_name = f"{zip_path.relative_to(PROJECT_ROOT)}::{member}"
    if not zip_path.exists():
        outputs.reject(source_code, str(zip_path), sheet, "", "event_normalized", "source file missing", {})
        return

    for row_number, row in iter_xlsx_rows_from_zip(zip_path, member, sheet):
        external_id = clean(row.get("ID_ECLO_TAMBO_VACA"))
        if not external_id:
            outputs.reject(source_code, file_name, sheet, row_number, "event_normalized", "missing ID_ECLO_TAMBO_VACA", row)
            continue
        canonical = add_map(map_records, source_code, external_id)

        calving_date = parse_date(row.get("FECHA"))
        if calving_date:
            event_id = outputs.write(
                "event_dogurma",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "dogurma_tarihi": calving_date,
                    "parite": row.get("lactancia"),
                    "source_table": "source_chiozza_sheet",
                    "source_candidate_id": row_number,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, calving_date, "calving", "event_dogurma", event_id, row)

        for service_column, service_label in [("primer_servicio", "first_service"), ("ult_servicio", "last_service")]:
            service_date = parse_date(row.get(service_column))
            if service_date:
                event_id = outputs.write(
                    "event_dollenme",
                    {
                        "source_code": source_code,
                        "file_name": file_name,
                        "sheet_name": sheet,
                        "source_row_number": row_number,
                        "external_animal_id": external_id,
                        "canonical_animal_id": canonical,
                        "dollenme_tarihi": service_date,
                        "servis_durumu": service_label,
                        "source_table": "source_chiozza_sheet",
                        "source_candidate_id": row_number,
                        "payload": payload_json(row),
                    },
                )
                add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, service_date, "insemination", "event_dollenme", event_id, row, text_value=service_label)

        diagnosis_date = parse_date(row.get("tacto_preniez"))
        if diagnosis_date and truthy(row.get("preniada")):
            add_normalized(
                outputs,
                source_code,
                file_name,
                sheet,
                row_number,
                external_id,
                diagnosis_date,
                "pregnancy_diagnosis_positive",
                "source_chiozza_sheet",
                row_number,
                row,
                text_value=row.get("preniada"),
            )

        lameness_date = parse_date(row.get("fecha_renguera"))
        if not lameness_date and truthy(row.get("renguera")):
            lameness_date = calving_date
        if truthy(row.get("renguera")) and lameness_date:
            raw_id = outputs.write(
                "raw_health",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "olay_tarihi": lameness_date,
                    "olay_tipi": "lameness",
                    "alt_tip": row.get("evento"),
                    "aciklama": "renguera flag/event from Chiozza dataset",
                    "raw_payload": payload_json(row),
                },
            )
            event_id = outputs.write(
                "event_health",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "olay_tarihi": lameness_date,
                    "olay_tipi": "lameness",
                    "alt_tip": row.get("evento"),
                    "aciklama": "renguera flag/event from Chiozza dataset",
                    "source_table": "raw_health",
                    "source_candidate_id": raw_id,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, lameness_date, "health_lameness", "event_health", event_id, row, text_value=row.get("evento"))
        elif truthy(row.get("renguera")):
            outputs.reject(source_code, file_name, sheet, row_number, "raw_health", "lameness flag present but no usable date", row)

        dry_date = parse_date(row.get("secado"))
        if dry_date:
            event_id = outputs.write(
                "event_kuru",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "kuru_baslangic_tarihi": dry_date,
                    "source_table": "source_chiozza_sheet",
                    "source_candidate_id": row_number,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, dry_date, "dryoff", "event_kuru", event_id, row)

        cull_date = parse_date(row.get("baja"))
        if clean(row.get("baja")) and cull_date:
            raw_id = outputs.write(
                "raw_culling",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "cikarilma_tarihi": cull_date,
                    "cikarilma_sebebi": "baja",
                    "raw_payload": payload_json(row),
                },
            )
            event_id = outputs.write(
                "event_suruden_cikarma",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "cikarilma_tarihi": cull_date,
                    "cikarilma_sebebi": "baja",
                    "source_table": "raw_culling",
                    "source_candidate_id": raw_id,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, cull_date, "cull", "event_suruden_cikarma", event_id, row, text_value="baja")
        elif clean(row.get("baja")):
            outputs.reject(source_code, file_name, sheet, row_number, "raw_culling", "baja is not a valid date", row)


def marcenat_date_fields(row: dict[str, Any], prefix: str) -> list[tuple[str, str]]:
    return [(key, parse_date(value)) for key, value in row.items() if key.startswith(prefix) and parse_date(value)]


def process_marcenat(outputs: CandidateOutputs, map_records: set[tuple[str, str, str]]) -> None:
    source_code = "marcenat_inrae"
    zip_path = EXTERNAL_DIR / source_code / "MarcenatINRAE.zip"
    member = "Original dataset.xlsx"
    sheet = "Feuil1"
    file_name = f"{zip_path.relative_to(PROJECT_ROOT)}::{member}"
    if not zip_path.exists():
        outputs.reject(source_code, str(zip_path), sheet, "", "event_normalized", "source file missing", {})
        return

    for row_number, row in iter_xlsx_rows_from_zip(zip_path, member, sheet):
        external_id = clean(row.get("Animal ID"))
        if not external_id:
            outputs.reject(source_code, file_name, sheet, row_number, "event_normalized", "missing Animal ID", row)
            continue
        canonical = add_map(map_records, source_code, external_id)

        for field_name, calving_date in marcenat_date_fields(row, "Date 1 er Mise Bas") + marcenat_date_fields(row, "Date Mise-bas"):
            event_id = outputs.write(
                "event_dogurma",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "dogurma_tarihi": calving_date,
                    "source_table": "source_marcenat_sheet",
                    "source_candidate_id": row_number,
                    "payload": payload_json({**row, "_event_field": field_name}),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, calving_date, "calving", "event_dogurma", event_id, {**row, "_event_field": field_name})

        for field_name in ["Date tarrissement", "Start of dry period"]:
            dry_date = parse_date(row.get(field_name))
            if dry_date:
                event_id = outputs.write(
                    "event_kuru",
                    {
                        "source_code": source_code,
                        "file_name": file_name,
                        "sheet_name": sheet,
                        "source_row_number": row_number,
                        "external_animal_id": external_id,
                        "canonical_animal_id": canonical,
                        "kuru_baslangic_tarihi": dry_date,
                        "source_table": "source_marcenat_sheet",
                        "source_candidate_id": row_number,
                        "payload": payload_json({**row, "_event_field": field_name}),
                    },
                )
                add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, dry_date, "dryoff", "event_kuru", event_id, {**row, "_event_field": field_name})

        cull_date = parse_date(row.get("Date Reforme"))
        if cull_date:
            raw_id = outputs.write(
                "raw_culling",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "cikarilma_tarihi": cull_date,
                    "cikarilma_sebebi": row.get("Reforme") or row.get("Cause Reforme"),
                    "raw_payload": payload_json(row),
                },
            )
            event_id = outputs.write(
                "event_suruden_cikarma",
                {
                    "source_code": source_code,
                    "file_name": file_name,
                    "sheet_name": sheet,
                    "source_row_number": row_number,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "cikarilma_tarihi": cull_date,
                    "cikarilma_sebebi": row.get("Reforme") or row.get("Cause Reforme"),
                    "source_table": "raw_culling",
                    "source_candidate_id": raw_id,
                    "payload": payload_json(row),
                },
            )
            add_normalized(outputs, source_code, file_name, sheet, row_number, external_id, cull_date, "cull", "event_suruden_cikarma", event_id, row, text_value=row.get("Reforme") or row.get("Cause Reforme"))


def write_map_records(map_records: set[tuple[str, str, str]]) -> None:
    with MAP_PATH.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=MAP_FIELDS)
        writer.writeheader()
        for source_code, external_id, canonical in sorted(map_records):
            writer.writerow(
                {
                    "source_code": source_code,
                    "external_animal_id": external_id,
                    "canonical_animal_id": canonical,
                    "notlar": "Generated by dry-run canonical ID rule.",
                }
            )


def main() -> int:
    outputs = CandidateOutputs()
    map_records: set[tuple[str, str, str]] = set()
    try:
        process_embryo(outputs, map_records)
        process_primiparous_daily(outputs, map_records)
        process_primiparous_events(outputs, map_records)
        process_chiozza(outputs, map_records)
        process_marcenat(outputs, map_records)
    finally:
        outputs.close()

    write_map_records(map_records)
    print(f"wrote {RAW_DIR.relative_to(PROJECT_ROOT)}")
    print(f"wrote {EVENT_DIR.relative_to(PROJECT_ROOT)}")
    print(f"wrote {EVENT_NORMALIZED_PATH.relative_to(PROJECT_ROOT)}")
    print(f"wrote {REJECTED_PATH.relative_to(PROJECT_ROOT)}")
    print(f"wrote {MAP_PATH.relative_to(PROJECT_ROOT)} ({len(map_records)} map rows)")
    for name, count in sorted(outputs.counts.items()):
        print(f"{name}: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
