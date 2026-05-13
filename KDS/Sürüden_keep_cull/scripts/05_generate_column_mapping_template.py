from __future__ import annotations

import csv
import re
import zipfile
from datetime import datetime
from io import BytesIO, TextIOWrapper
from pathlib import Path
from typing import BinaryIO

import yaml
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXTERNAL_DIR = PROJECT_ROOT / "data" / "external"
OUTPUT_PATH = PROJECT_ROOT / "reports" / "column_mapping_template.csv"
RULES_PATH = PROJECT_ROOT / "config" / "canonical_event_rules.yaml"
SUPPORTED_SUFFIXES = {".csv", ".tsv", ".tab", ".txt", ".xlsx", ".zip"}

FIELDNAMES = [
    "source_code",
    "file_name",
    "sheet_name",
    "source_column",
    "detected_type",
    "suggested_target_table",
    "suggested_target_column",
    "confidence",
    "notes",
]


def normalize_key(value: str) -> str:
    return re.sub(r"\s+", " ", str(value).strip()).lower()


def load_rules() -> tuple[dict[tuple[str, str, str], dict], dict[tuple[str, str], dict]]:
    if not RULES_PATH.exists():
        return {}, {}
    with RULES_PATH.open("r", encoding="utf-8") as file:
        rules = yaml.safe_load(file) or {}

    overrides: dict[tuple[str, str, str], dict] = {}
    sheet_configs: dict[tuple[str, str], dict] = {}
    for source_code, source_config in (rules.get("sources") or {}).items():
        for sheet_name, sheet_config in (source_config.get("sheets") or {}).items():
            sheet_configs[(source_code, normalize_key(sheet_name))] = sheet_config or {}
            for column_name, column_rule in (sheet_config.get("columns") or {}).items():
                overrides[(source_code, normalize_key(sheet_name), normalize_key(column_name))] = column_rule or {}
    return overrides, sheet_configs


RULE_OVERRIDES, SHEET_CONFIGS = load_rules()


def source_rule_override(source_code: str, sheet_name: str, column_name: str) -> dict | None:
    exact = RULE_OVERRIDES.get((source_code, normalize_key(sheet_name), normalize_key(column_name)))
    if exact is not None:
        return exact

    # Some files are inspected with an empty sheet name, while the rules may use
    # a filename-like pseudo sheet for CSV/manual notes. Fall back to any exact
    # column rule defined for the same source if it is unique.
    candidates = [
        rule
        for (rule_source, _rule_sheet, rule_column), rule in RULE_OVERRIDES.items()
        if rule_source == source_code and rule_column == normalize_key(column_name)
    ]
    if len(candidates) == 1:
        return candidates[0]
    return None


def sheet_config(source_code: str, sheet_name: str) -> dict:
    exact = SHEET_CONFIGS.get((source_code, normalize_key(sheet_name)))
    if exact is not None:
        return exact
    candidates = [config for (rule_source, _rule_sheet), config in SHEET_CONFIGS.items() if rule_source == source_code]
    return candidates[0] if len(candidates) == 1 else {}


def should_suppress_heuristic(source_code: str, sheet_name: str) -> tuple[bool, str]:
    config = sheet_config(source_code, sheet_name)
    row_action = str(config.get("row_action") or "")
    policy = str(config.get("heuristic_policy") or "")
    if policy == "rules_only":
        reason = str(config.get("skip_reason") or "No source-specific canonical rule for this column.")
        return True, reason
    if row_action.startswith("skip_"):
        reason = str(config.get("skip_reason") or "Sheet skipped by canonical rules.")
        return True, reason
    return False, ""


def source_code_for_path(path: Path) -> str:
    try:
        rel = path.relative_to(EXTERNAL_DIR)
        return rel.parts[0] if rel.parts else ""
    except ValueError:
        return ""


def csv_dialect(sample: str, path_or_name: str) -> csv.Dialect:
    suffix = Path(path_or_name).suffix.lower()
    if suffix in {".tsv", ".tab"}:
        return csv.get_dialect("excel-tab")
    try:
        return csv.Sniffer().sniff(sample)
    except csv.Error:
        return csv.get_dialect("excel")


def read_csv_head(file: TextIOWrapper, path_or_name: str, max_rows: int = 21) -> tuple[list[str], list[list[str]]]:
    sample = file.read(4096)
    file.seek(0)
    dialect = csv_dialect(sample, path_or_name)
    reader = csv.reader(file, dialect)
    rows = []
    for index, row in enumerate(reader):
        rows.append(row)
        if index >= max_rows - 1:
            break
    if not rows:
        return [], []
    return [str(value) for value in rows[0]], rows[1:]


def is_numeric(value: str) -> bool:
    try:
        float(value.replace(",", ""))
        return True
    except ValueError:
        return False


def is_date(value: str) -> bool:
    value = value.strip()
    if not value:
        return False
    candidates = [
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d.%m.%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%Y-%m-%d %H:%M:%S",
    ]
    for fmt in candidates:
        try:
            datetime.strptime(value, fmt)
            return True
        except ValueError:
            continue
    try:
        datetime.fromisoformat(value)
        return True
    except ValueError:
        return False


def detect_type(values: list[object]) -> str:
    cleaned = [str(value).strip() for value in values if value is not None and str(value).strip() != ""]
    if not cleaned:
        return "unknown"

    numeric_ratio = sum(1 for value in cleaned if is_numeric(value)) / len(cleaned)
    if numeric_ratio >= 0.8:
        return "numeric"

    date_ratio = sum(1 for value in cleaned if is_date(value)) / len(cleaned)
    if date_ratio >= 0.8:
        return "date"

    lowered = {value.lower() for value in cleaned}
    if lowered <= {"true", "false", "yes", "no", "y", "n", "0", "1", "e", "h"}:
        return "boolean_or_flag"

    return "text"


def suggest_mapping(column_name: str, detected_type: str) -> tuple[str, str, str, str]:
    col = column_name.lower().strip()
    compact = re.sub(r"[^a-z0-9]+", "_", col).strip("_")

    if re.search(r"(^|_)(cowid|cow_id|animal_id|animal|cow|id)(_|$)", compact):
        return "map_external_inek", "external_animal_id", "0.70", "Animal/cow/id heuristic; confirm unique ID semantics."

    if any(token in compact for token in ["milk", "yield", "leche", "my", "sut"]):
        return "raw_milk", "sut_miktar_kg", "0.65", "Milk/yield heuristic; confirm unit and aggregation window."

    if any(token in compact for token in ["test_date", "date", "day", "fecha", "tarih", "jour"]):
        return "event_normalized", "event_date", "0.55", "Date/day heuristic; choose source-specific target date after manual review."

    if any(token in compact for token in ["cull", "removal", "exit", "disposal", "baja", "reforme"]):
        if any(token in compact for token in ["reason", "cause", "type"]):
            return "raw_culling", "cikarilma_sebebi", "0.70", "Culling reason heuristic."
        if detected_type == "date":
            return "raw_culling", "cikarilma_tarihi", "0.70", "Culling/removal date heuristic."
        return "raw_culling", "raw_payload", "0.50", "Culling heuristic; choose date/reason/value after manual review."

    if any(token in compact for token in ["dry", "dryoff", "secado", "tarrissement"]):
        return "event_kuru", "kuru_baslangic_tarihi" if detected_type == "date" else "source_row_id", "0.55", "Dry-off heuristic; confirm event date semantics."

    if any(token in compact for token in ["mastitis", "lameness", "ketosis", "metritis", "renguera", "boiterie", "mam"]):
        return "raw_health", "olay_tipi", "0.75", "Health event keyword heuristic."

    if any(token in compact for token in ["calving", "fparto", "birth", "mise_bas", "velage", "parto"]):
        target_col = "dogurma_tarihi" if detected_type == "date" else "raw_payload"
        target_table = "event_dogurma" if detected_type == "date" else "raw_repro_loss"
        return target_table, target_col, "0.65", "Calving/birth heuristic; confirm whether this is cow birth or calving event."

    if any(token in compact for token in ["service", "servicio", "servicios", "insemination", "inseminations", "ia", "sn"]):
        target_col = "dollenme_tarihi" if detected_type == "date" else "dollenme_sayisi"
        return "event_dollenme", target_col, "0.65", "Service/insemination heuristic."

    if any(token in compact for token in ["heat", "chaleur"]):
        return "raw_repro_loss", "raw_payload", "0.45", "Heat/reproduction heuristic; map manually before raw load."

    if "scc" in compact:
        return "raw_milk", "scc", "0.80", "SCC heuristic."

    if "cell" in compact:
        return "raw_milk", "scc", "0.55", "Cell count heuristic; confirm SCC scale and unit."

    if "conductivity" in compact:
        return "raw_milk", "conductivity", "0.80", "Conductivity heuristic."

    if any(token in compact for token in ["fat", "yag", "tb", "mf"]):
        return "raw_milk", "yag_orani", "0.70", "Fat percentage heuristic; confirm unit."

    if any(token in compact for token in ["protein", "tp", "mp"]):
        return "raw_milk", "protein_orani", "0.70", "Protein percentage heuristic; confirm unit."

    if any(token in compact for token in ["gestation", "pregnancy", "preniez", "preniada", "perdida", "fechaperd"]):
        return "raw_repro_loss", "raw_payload", "0.50", "Reproduction/pregnancy heuristic; map exact meaning manually."

    return "", "", "0.00", "TODO / manual inspection required."


def suggest_economic_mapping(source_code: str, column_name: str, detected_type: str) -> tuple[str, str, str, str] | None:
    col = column_name.lower().strip()
    compact = re.sub(r"[^a-z0-9]+", "_", col).strip("_")

    if source_code == "usda_nass_milk_price":
        if compact in {"year", "period"}:
            return "feat_reward_parametre", "gecerlilik_baslangic", "0.50", "NASS year/period define monthly validity; combine manually into a date."
        if compact == "value":
            return "feat_reward_parametre", "sut_fiyati", "0.55", "NASS milk price value; verify Excel did not coerce decimals to dates and convert $/CWT before insert."
        if compact in {"data_item", "commodity", "state", "geo_level", "domain", "domain_category"}:
            return "feat_reward_parametre", "aciklama", "0.45", "NASS metadata column; preserve in scenario notes or source payload."

    if source_code == "usda_ers_milk_cost":
        if compact == "year":
            return "feat_reward_parametre", "gecerlilik_baslangic", "0.50", "ERS annual estimate year; use first day of year unless another convention is approved."
        if compact == "value":
            return "feat_reward_parametre", "", "0.40", "ERS economic value; target parameter depends on Item and Units, manual unit conversion required."
        if compact in {"category", "item", "item2", "units", "size", "region", "country", "survey_base_year"}:
            return "feat_reward_parametre", "aciklama", "0.45", "ERS metadata column; preserve in scenario notes or source payload."

    return None


def rows_from_columns(
    source_code: str,
    file_name: str,
    sheet_name: str,
    columns: list[str],
    sample_rows: list[list[object]],
) -> list[dict]:
    output: list[dict] = []
    width = len(columns)
    for index, column_name in enumerate(columns):
        values = [row[index] for row in sample_rows if index < len(row)]
        detected_type = detect_type(values)
        rule = source_rule_override(source_code, sheet_name, column_name)
        if rule is not None:
            target_table = str(rule.get("target_table") or "")
            target_column = str(rule.get("target_column") or "")
            confidence = f"{float(rule.get('confidence', 0.0)):.2f}"
            notes = str(rule.get("notes") or "Source-specific canonical rule.")
            if rule.get("semantic"):
                notes = f"{rule.get('semantic')}. {notes}"
        else:
            suppress_heuristic, suppress_reason = should_suppress_heuristic(source_code, sheet_name)
            economic_suggestion = None if suppress_heuristic else suggest_economic_mapping(source_code, column_name, detected_type)
            if economic_suggestion is not None:
                target_table, target_column, confidence, notes = economic_suggestion
            elif suppress_heuristic:
                target_table = ""
                target_column = ""
                confidence = "0.00"
                notes = f"TODO / manual inspection required. {suppress_reason}"
            else:
                target_table, target_column, confidence, notes = suggest_mapping(column_name, detected_type)
        output.append(
            {
                "source_code": source_code,
                "file_name": file_name,
                "sheet_name": sheet_name,
                "source_column": column_name,
                "detected_type": detected_type,
                "suggested_target_table": target_table,
                "suggested_target_column": target_column,
                "confidence": confidence,
                "notes": notes,
            }
        )
    if width == 0:
        output.append(
            {
                "source_code": source_code,
                "file_name": file_name,
                "sheet_name": sheet_name,
                "source_column": "",
                "detected_type": "unknown",
                "suggested_target_table": "",
                "suggested_target_column": "",
                "confidence": "0.00",
                "notes": "TODO / manual inspection required. No columns detected.",
            }
        )
    return output


def inspect_csv_path(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8", errors="replace", newline="") as file:
        columns, sample_rows = read_csv_head(file, path.name)
    return rows_from_columns(
        source_code=source_code_for_path(path),
        file_name=str(path.relative_to(PROJECT_ROOT)),
        sheet_name="",
        columns=columns,
        sample_rows=sample_rows,
    )


def inspect_csv_zip_member(zip_file: zipfile.ZipFile, zip_path: Path, member: str) -> list[dict]:
    with zip_file.open(member) as raw:
        text_stream = TextIOWrapper(raw, encoding="utf-8", errors="replace", newline="")
        columns, sample_rows = read_csv_head(text_stream, member)
    return rows_from_columns(
        source_code=source_code_for_path(zip_path),
        file_name=f"{zip_path.relative_to(PROJECT_ROOT)}::{member}",
        sheet_name="",
        columns=columns,
        sample_rows=sample_rows,
    )


def workbook_rows(stream: BinaryIO | BytesIO, source_code: str, file_name: str) -> list[dict]:
    workbook = load_workbook(stream, read_only=True, data_only=True)
    output: list[dict] = []
    for worksheet in workbook.worksheets:
        raw_rows = list(worksheet.iter_rows(min_row=1, max_row=21, values_only=True))
        if not raw_rows:
            continue
        columns = [str(value) if value is not None else "" for value in raw_rows[0]]
        sample_rows = [list(row) for row in raw_rows[1:]]
        output.extend(rows_from_columns(source_code, file_name, worksheet.title, columns, sample_rows))
    return output


def inspect_xlsx_path(path: Path) -> list[dict]:
    with path.open("rb") as stream:
        return workbook_rows(stream, source_code_for_path(path), str(path.relative_to(PROJECT_ROOT)))


def inspect_zip(path: Path) -> list[dict]:
    rows: list[dict] = []
    with zipfile.ZipFile(path) as zip_file:
        for member in zip_file.namelist():
            if member.endswith("/"):
                continue
            suffix = Path(member).suffix.lower()
            try:
                if suffix in {".csv", ".tsv", ".tab", ".txt"}:
                    rows.extend(inspect_csv_zip_member(zip_file, path, member))
                elif suffix == ".xlsx":
                    with zip_file.open(member) as raw:
                        rows.extend(
                            workbook_rows(
                                BytesIO(raw.read()),
                                source_code_for_path(path),
                                f"{path.relative_to(PROJECT_ROOT)}::{member}",
                            )
                        )
                else:
                    rows.append(
                        {
                            "source_code": source_code_for_path(path),
                            "file_name": f"{path.relative_to(PROJECT_ROOT)}::{member}",
                            "sheet_name": "",
                            "source_column": "",
                            "detected_type": "unknown",
                            "suggested_target_table": "",
                            "suggested_target_column": "",
                            "confidence": "0.00",
                            "notes": f"TODO / manual inspection required. Unsupported file type in ZIP: {suffix or 'unknown'}",
                        }
                    )
            except Exception as exc:  # noqa: BLE001
                rows.append(
                    {
                        "source_code": source_code_for_path(path),
                        "file_name": f"{path.relative_to(PROJECT_ROOT)}::{member}",
                        "sheet_name": "",
                        "source_column": "",
                        "detected_type": "unknown",
                        "suggested_target_table": "",
                        "suggested_target_column": "",
                        "confidence": "0.00",
                        "notes": f"TODO / manual inspection required. Error: {exc}",
                    }
                )
    return rows


def supported_files() -> list[Path]:
    return sorted(
        path
        for path in EXTERNAL_DIR.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )


def main() -> int:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    rows: list[dict] = []

    for path in supported_files():
        try:
            suffix = path.suffix.lower()
            if suffix in {".csv", ".tsv", ".tab", ".txt"}:
                rows.extend(inspect_csv_path(path))
            elif suffix == ".xlsx":
                rows.extend(inspect_xlsx_path(path))
            elif suffix == ".zip":
                rows.extend(inspect_zip(path))
        except Exception as exc:  # noqa: BLE001
            rows.append(
                {
                    "source_code": source_code_for_path(path),
                    "file_name": str(path.relative_to(PROJECT_ROOT)),
                    "sheet_name": "",
                    "source_column": "",
                    "detected_type": "unknown",
                    "suggested_target_table": "",
                    "suggested_target_column": "",
                    "confidence": "0.00",
                    "notes": f"TODO / manual inspection required. Error: {exc}",
                }
            )

    with OUTPUT_PATH.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)

    print(f"wrote {OUTPUT_PATH.relative_to(PROJECT_ROOT)} ({len(rows)} rows)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
