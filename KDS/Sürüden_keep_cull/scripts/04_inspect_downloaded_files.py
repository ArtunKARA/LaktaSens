from __future__ import annotations

import csv
import zipfile
from io import BytesIO, TextIOWrapper
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXTERNAL_DIR = PROJECT_ROOT / "data" / "external"
REPORT_PATH = PROJECT_ROOT / "reports" / "file_inspection_report.md"
SUPPORTED_SUFFIXES = {".csv", ".tsv", ".tab", ".txt", ".xlsx", ".zip"}


def markdown_code(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("`", "'")


def normalized_rows(rows: list[list[object]], width: int) -> list[list[str]]:
    output: list[list[str]] = []
    for row in rows:
        values = [markdown_code(value) for value in row]
        if len(values) < width:
            values.extend([""] * (width - len(values)))
        output.append(values[:width])
    return output


def sample_to_text(rows: list[list[object]], columns: list[str]) -> str:
    if not rows:
        return "_No sample rows._"
    width = max(len(columns), max((len(row) for row in rows), default=0))
    headers = normalized_rows([columns], width)[0]
    body = normalized_rows(rows, width)
    lines = ["\t".join(headers)]
    lines.extend("\t".join(row) for row in body[:5])
    return "```text\n" + "\n".join(lines) + "\n```"


def csv_dialect(sample: str, path_or_name: str) -> csv.Dialect:
    suffix = Path(path_or_name).suffix.lower()
    if suffix in {".tsv", ".tab"}:
        return csv.get_dialect("excel-tab")
    try:
        return csv.Sniffer().sniff(sample)
    except csv.Error:
        return csv.get_dialect("excel")


def count_text_rows(path: Path) -> int | None:
    try:
        with path.open("r", encoding="utf-8", errors="replace", newline="") as file:
            line_count = sum(1 for _ in file)
        return max(line_count - 1, 0)
    except Exception:  # noqa: BLE001
        return None


def count_zip_text_rows(zip_file: zipfile.ZipFile, member: str) -> int | None:
    try:
        with zip_file.open(member) as raw:
            line_count = sum(1 for _ in raw)
        return max(line_count - 1, 0)
    except Exception:  # noqa: BLE001
        return None


def read_csv_head_from_text_stream(file: TextIOWrapper, path_or_name: str) -> tuple[list[str], list[list[str]]]:
    sample = file.read(4096)
    file.seek(0)
    dialect = csv_dialect(sample, path_or_name)
    reader = csv.reader(file, dialect)
    rows = []
    for index, row in enumerate(reader):
        rows.append(row)
        if index >= 5:
            break
    if not rows:
        return [], []
    return [str(value) for value in rows[0]], rows[1:6]


def inspect_csv_path(path: Path, display_name: str) -> list[str]:
    with path.open("r", encoding="utf-8", errors="replace", newline="") as file:
        columns, sample_rows = read_csv_head_from_text_stream(file, path.name)
    return render_file_section(
        file_name=display_name,
        file_format=path.suffix.lower().lstrip("."),
        sheet_name="",
        row_count=count_text_rows(path),
        columns=columns,
        sample_rows=sample_rows,
    )


def inspect_csv_zip_member(zip_file: zipfile.ZipFile, member: str, display_name: str) -> list[str]:
    with zip_file.open(member) as raw:
        text_stream = TextIOWrapper(raw, encoding="utf-8", errors="replace", newline="")
        columns, sample_rows = read_csv_head_from_text_stream(text_stream, member)
    return render_file_section(
        file_name=display_name,
        file_format=Path(member).suffix.lower().lstrip("."),
        sheet_name="",
        row_count=count_zip_text_rows(zip_file, member),
        columns=columns,
        sample_rows=sample_rows,
    )


def inspect_xlsx_stream(stream: BinaryIO | BytesIO, display_name: str) -> list[str]:
    workbook = load_workbook(stream, read_only=True, data_only=True)
    sections: list[str] = []
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(min_row=1, max_row=6, values_only=True))
        header = [markdown_code(value) for value in rows[0]] if rows else []
        sample_rows = [list(row) for row in rows[1:6]] if len(rows) > 1 else []
        sections.extend(
            render_file_section(
                file_name=display_name,
                file_format="xlsx",
                sheet_name=worksheet.title,
                row_count=max((worksheet.max_row or 1) - 1, 0),
                columns=header,
                sample_rows=sample_rows,
            )
        )
    return sections


def inspect_xlsx_path(path: Path, display_name: str) -> list[str]:
    with path.open("rb") as stream:
        return inspect_xlsx_stream(stream, display_name)


def render_file_section(
    file_name: str,
    file_format: str,
    sheet_name: str,
    row_count: int | None,
    columns: list[str],
    sample_rows: list[list[object]] | None = None,
) -> list[str]:
    lines = [
        f"## {file_name}",
        "",
        f"- Format: `{file_format}`",
        f"- Sheet name: `{sheet_name}`" if sheet_name else "- Sheet name: ",
        f"- Row count: `{row_count if row_count is not None else 'unknown'}`",
        f"- Columns: {', '.join(f'`{markdown_code(col)}`' for col in columns) if columns else 'TODO / manual inspection required'}",
        "",
        "First 5 rows:",
        "",
        sample_to_text(sample_rows or [], columns),
        "",
    ]
    return lines


def inspect_zip(path: Path) -> list[str]:
    sections = [f"## {path.relative_to(PROJECT_ROOT)}", "", "- Format: `zip`", ""]
    with zipfile.ZipFile(path) as zip_file:
        members = [name for name in zip_file.namelist() if not name.endswith("/")]
        sections.append(f"- Member count: `{len(members)}`")
        sections.append("")
        for member in members:
            suffix = Path(member).suffix.lower()
            display_name = f"{path.relative_to(PROJECT_ROOT)}::{member}"
            try:
                if suffix in {".csv", ".tsv", ".tab", ".txt"}:
                    sections.extend(inspect_csv_zip_member(zip_file, member, display_name))
                elif suffix == ".xlsx":
                    with zip_file.open(member) as raw:
                        sections.extend(inspect_xlsx_stream(BytesIO(raw.read()), display_name))
                else:
                    sections.extend(
                        [
                            f"### {display_name}",
                            "",
                            f"- Format: `{suffix.lstrip('.') or 'unknown'}`",
                            "- Status: skipped; unsupported file type inside ZIP.",
                            "",
                        ]
                    )
            except Exception as exc:  # noqa: BLE001
                sections.extend(
                    [
                        f"### {display_name}",
                        "",
                        "- Status: TODO / manual inspection required",
                        f"- Error: `{exc}`",
                        "",
                    ]
                )
    return sections


def supported_files() -> list[Path]:
    return sorted(
        path
        for path in EXTERNAL_DIR.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )


def main() -> int:
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# File Inspection Report",
        "",
        "Generated from files under `data/external/`.",
        "",
    ]

    files = supported_files()
    if not files:
        lines.extend(
            [
                "No supported CSV, TSV/TAB, XLSX or ZIP files found.",
                "",
                "Downloaded PDFs are kept under `data/external/` but are not schema-inspected by this script.",
                "",
                "Run `python scripts/03_download_sources.py` or place manual downloads under `data/external/<source_code>/`.",
                "",
            ]
        )
    else:
        for path in files:
            display_name = str(path.relative_to(PROJECT_ROOT))
            try:
                suffix = path.suffix.lower()
                if suffix in {".csv", ".tsv", ".tab", ".txt"}:
                    lines.extend(inspect_csv_path(path, display_name))
                elif suffix == ".xlsx":
                    lines.extend(inspect_xlsx_path(path, display_name))
                elif suffix == ".zip":
                    lines.extend(inspect_zip(path))
            except Exception as exc:  # noqa: BLE001
                lines.extend(
                    [
                        f"## {display_name}",
                        "",
                        "- Status: TODO / manual inspection required",
                        f"- Error: `{exc}`",
                        "",
                    ]
                )

    REPORT_PATH.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    print(f"wrote {REPORT_PATH.relative_to(PROJECT_ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
