from __future__ import annotations

import csv
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
NASS_FILE = PROJECT_ROOT / "data" / "external" / "usda_nass_milk_price" / "USDA NASS.xlsx"
ERS_STATE_FILE = PROJECT_ROOT / "data" / "external" / "usda_ers_milk_cost" / "MilkState2021base.csv"
OUTPUT_CSV = PROJECT_ROOT / "data" / "processed" / "reward_baseline_us_2025.csv"
OUTPUT_MD = PROJECT_ROOT / "reports" / "reward_baseline_us_2025_review.md"
OUTPUT_SQL = PROJECT_ROOT / "sql" / "08_seed_reward_parametre_us_2025.sql"

SCENARIO_CODE = "baseline_us_2025_mixed_source_v1"
LEGACY_SCENARIO_CODE = "baseline_external_pending"
VALID_FROM = "2025-01-01"
VALID_TO = "2025-12-31"
CWT_TO_KG = Decimal("45.359237")
MONTHS = {"JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"}

# These are explicit prototype assumptions because the downloaded local files do
# not expose parseable, farm-specific values for these reward terms.
SYNTHETIC_REPLACEMENT_COST_USD = Decimal("2500.00")
SYNTHETIC_CULL_SALE_VALUE_USD = Decimal("1200.00")
SYNTHETIC_TREATMENT_COST_USD = Decimal("150.00")
SYNTHETIC_FAILED_INSEMINATION_COST_USD = Decimal("35.00")
SYNTHETIC_LOW_YIELD_KG_PER_DAY = Decimal("5.00")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def decimal_value(value: Any) -> Decimal:
    text = clean_text(value).replace(",", "")
    if not text:
        raise ValueError("empty numeric value")
    return Decimal(text)


def q(value: Decimal, places: str) -> Decimal:
    return value.quantize(Decimal(places), rounding=ROUND_HALF_UP)


def sql_text(value: str) -> str:
    return "$$" + value.replace("$$", "$ $") + "$$"


def read_nass_2025_monthly_prices() -> list[dict[str, Any]]:
    workbook = load_workbook(NASS_FILE, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        records = worksheet.iter_rows(values_only=True)
        try:
            headers = [clean_text(value) for value in next(records)]
        except StopIteration:
            continue
        for record in records:
            row = dict(zip(headers, record))
            year = clean_text(row.get("Year"))
            period = clean_text(row.get("Period")).upper()
            state = clean_text(row.get("State")).upper()
            data_item = clean_text(row.get("Data Item")).upper()
            if year != "2025":
                continue
            if period not in MONTHS:
                continue
            if state and state != "US TOTAL":
                continue
            if "PRICE RECEIVED" not in data_item:
                continue
            rows.append(
                {
                    "period": period,
                    "value_per_cwt": decimal_value(row.get("Value")),
                    "source_row": row,
                }
            )
    if not rows:
        raise RuntimeError(f"No 2025 monthly NASS milk price rows found in {NASS_FILE}")
    return rows


def read_ers_rows() -> list[dict[str, str]]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with ERS_STATE_FILE.open("r", encoding=encoding, newline="") as file:
                return list(csv.DictReader(file))
        except UnicodeDecodeError:
            continue
    raise RuntimeError(f"Could not decode {ERS_STATE_FILE}")


def ers_value(rows: list[dict[str, str]], item: str) -> Decimal:
    matches = [
        row
        for row in rows
        if clean_text(row.get("Year")) == "2025"
        and clean_text(row.get("Region")).lower() == "u.s. total"
        and clean_text(row.get("Size")).lower() == "no specific size"
        and clean_text(row.get("Item")) == item
    ]
    if not matches:
        raise RuntimeError(f"ERS 2025 U.S. total row not found for item: {item}")
    return decimal_value(matches[0].get("Value"))


def build_values() -> dict[str, Decimal | str | int]:
    nass_rows = read_nass_2025_monthly_prices()
    nass_avg_per_cwt = sum(row["value_per_cwt"] for row in nass_rows) / Decimal(len(nass_rows))
    milk_price_usd_per_kg = nass_avg_per_cwt / CWT_TO_KG

    ers_rows = read_ers_rows()
    output_per_cow_lb_year = ers_value(ers_rows, "Output per cow")
    output_per_cow_cwt_year = output_per_cow_lb_year / Decimal("100")
    operating_cost_per_cwt = ers_value(ers_rows, "Total, operating costs")
    total_costs_listed_per_cwt = ers_value(ers_rows, "Total, costs listed")
    feed_cost_per_cwt = ers_value(ers_rows, "Total, feed costs")
    milk_sold_per_cwt = ers_value(ers_rows, "Milk sold")
    milk_cows_head_per_farm = ers_value(ers_rows, "Milk cows")

    daily_operating_cost = operating_cost_per_cwt * output_per_cow_cwt_year / Decimal("365")
    daily_total_costs_listed = total_costs_listed_per_cwt * output_per_cow_cwt_year / Decimal("365")
    daily_feed_cost = feed_cost_per_cwt * output_per_cow_cwt_year / Decimal("365")
    low_yield_penalty = milk_price_usd_per_kg * SYNTHETIC_LOW_YIELD_KG_PER_DAY

    return {
        "nass_month_count": len(nass_rows),
        "nass_avg_per_cwt": q(nass_avg_per_cwt, "0.001"),
        "milk_price_usd_per_kg": q(milk_price_usd_per_kg, "0.0001"),
        "ers_milk_sold_per_cwt": q(milk_sold_per_cwt, "0.01"),
        "ers_output_per_cow_lb_year": q(output_per_cow_lb_year, "0.01"),
        "ers_output_per_cow_cwt_year": q(output_per_cow_cwt_year, "0.01"),
        "ers_operating_cost_per_cwt": q(operating_cost_per_cwt, "0.01"),
        "ers_total_costs_listed_per_cwt": q(total_costs_listed_per_cwt, "0.01"),
        "ers_feed_cost_per_cwt": q(feed_cost_per_cwt, "0.01"),
        "ers_milk_cows_head_per_farm": q(milk_cows_head_per_farm, "0.01"),
        "daily_operating_cost": q(daily_operating_cost, "0.0001"),
        "daily_total_costs_listed": q(daily_total_costs_listed, "0.0001"),
        "daily_feed_cost": q(daily_feed_cost, "0.0001"),
        "replacement_cost": SYNTHETIC_REPLACEMENT_COST_USD,
        "cull_sale_value": SYNTHETIC_CULL_SALE_VALUE_USD,
        "treatment_cost": SYNTHETIC_TREATMENT_COST_USD,
        "failed_insemination_cost": SYNTHETIC_FAILED_INSEMINATION_COST_USD,
        "low_yield_kg_per_day": SYNTHETIC_LOW_YIELD_KG_PER_DAY,
        "low_yield_penalty": q(low_yield_penalty, "0.01"),
    }


def baseline_notes(values: dict[str, Decimal | str | int]) -> str:
    return (
        "US 2025 prototype baseline. Source-derived: NASS 2025 US TOTAL monthly average "
        f"{values['nass_avg_per_cwt']} USD/cwt converted to {values['milk_price_usd_per_kg']} USD/kg; "
        f"ERS operating cost {values['ers_operating_cost_per_cwt']} USD/cwt with "
        f"{values['ers_output_per_cow_cwt_year']} cwt/cow/year converted to "
        f"{values['daily_operating_cost']} USD/cow/day. Synthetic assumptions: "
        f"replacement_maliyeti={values['replacement_cost']} USD/head, "
        f"cull_satis_degeri={values['cull_sale_value']} USD/head, "
        f"tedavi_maliyeti={values['treatment_cost']} USD/event, "
        f"basarisiz_tohumlama_maliyeti={values['failed_insemination_cost']} USD/service, "
        f"dusuk_verim_cezasi={values['low_yield_kg_per_day']} kg/day * milk price = "
        f"{values['low_yield_penalty']} USD/day. Adjust before production use."
    )


def write_baseline_csv(values: dict[str, Decimal | str | int]) -> None:
    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    rows = [
        {
            "parameter": "sut_fiyati",
            "value": values["milk_price_usd_per_kg"],
            "unit": "USD/kg",
            "value_source": "source_derived",
            "source_file": str(NASS_FILE.relative_to(PROJECT_ROOT)),
            "formula": "avg(2025 monthly US TOTAL NASS milk price USD/cwt) / 45.359237 kg/cwt",
            "notes": f"NASS month count={values['nass_month_count']}; avg={values['nass_avg_per_cwt']} USD/cwt.",
        },
        {
            "parameter": "gunluk_bakim_maliyeti",
            "value": values["daily_operating_cost"],
            "unit": "USD/cow/day",
            "value_source": "source_derived",
            "source_file": str(ERS_STATE_FILE.relative_to(PROJECT_ROOT)),
            "formula": "ERS operating cost USD/cwt * ERS output per cow cwt/year / 365",
            "notes": "Uses operating costs, not total listed costs, to avoid double counting overhead in the first prototype.",
        },
        {
            "parameter": "reference_total_costs_listed_daily",
            "value": values["daily_total_costs_listed"],
            "unit": "USD/cow/day",
            "value_source": "source_derived_reference_only",
            "source_file": str(ERS_STATE_FILE.relative_to(PROJECT_ROOT)),
            "formula": "ERS total costs listed USD/cwt * ERS output per cow cwt/year / 365",
            "notes": "Reference value only; not inserted as gunluk_bakim_maliyeti.",
        },
        {
            "parameter": "reference_feed_cost_daily",
            "value": values["daily_feed_cost"],
            "unit": "USD/cow/day",
            "value_source": "source_derived_reference_only",
            "source_file": str(ERS_STATE_FILE.relative_to(PROJECT_ROOT)),
            "formula": "ERS feed cost USD/cwt * ERS output per cow cwt/year / 365",
            "notes": "Reference value only.",
        },
        {
            "parameter": "replacement_maliyeti",
            "value": values["replacement_cost"],
            "unit": "USD/head",
            "value_source": "synthetic_assumption",
            "source_file": "",
            "formula": "fixed prototype assumption",
            "notes": "Downloaded Iowa State PDF was not parsed into a reliable table in this phase.",
        },
        {
            "parameter": "cull_satis_degeri",
            "value": values["cull_sale_value"],
            "unit": "USD/head",
            "value_source": "synthetic_assumption",
            "source_file": "",
            "formula": "fixed prototype assumption",
            "notes": "Downloaded AMS PDF was not parsed into a reliable market-class table in this phase.",
        },
        {
            "parameter": "tedavi_maliyeti",
            "value": values["treatment_cost"],
            "unit": "USD/event",
            "value_source": "synthetic_assumption",
            "source_file": "",
            "formula": "fixed prototype assumption",
            "notes": "Health treatment cost is farm/protocol-specific.",
        },
        {
            "parameter": "basarisiz_tohumlama_maliyeti",
            "value": values["failed_insemination_cost"],
            "unit": "USD/service",
            "value_source": "synthetic_assumption",
            "source_file": "",
            "formula": "fixed prototype assumption",
            "notes": "Insemination cost is farm/protocol-specific.",
        },
        {
            "parameter": "dusuk_verim_cezasi",
            "value": values["low_yield_penalty"],
            "unit": "USD/day",
            "value_source": "mixed_source_and_synthetic",
            "source_file": str(NASS_FILE.relative_to(PROJECT_ROOT)),
            "formula": "NASS milk price USD/kg * synthetic 5 kg/day low-yield gap",
            "notes": "Penalty scale should be tuned with farm policy.",
        },
    ]
    fieldnames = ["parameter", "value", "unit", "value_source", "source_file", "formula", "notes"]
    with OUTPUT_CSV.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_review(values: dict[str, Decimal | str | int]) -> None:
    OUTPUT_MD.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Reward Baseline US 2025 Review",
        "",
        f"- DB scenario: `{SCENARIO_CODE}`",
        f"- Validity: `{VALID_FROM}` to `{VALID_TO}`",
        "- This is a prototype baseline, not a production financial recommendation.",
        "- Source-derived values are separated from synthetic assumptions below.",
        "",
        "## DB Parameters",
        "",
        "| Parametre | Deger | Birim | Kaynak durumu |",
        "|---|---:|---|---|",
        f"| `sut_fiyati` | {values['milk_price_usd_per_kg']} | USD/kg | NASS source-derived |",
        f"| `gunluk_bakim_maliyeti` | {values['daily_operating_cost']} | USD/cow/day | ERS source-derived |",
        f"| `replacement_maliyeti` | {values['replacement_cost']} | USD/head | synthetic assumption |",
        f"| `cull_satis_degeri` | {values['cull_sale_value']} | USD/head | synthetic assumption |",
        f"| `tedavi_maliyeti` | {values['treatment_cost']} | USD/event | synthetic assumption |",
        f"| `basarisiz_tohumlama_maliyeti` | {values['failed_insemination_cost']} | USD/service | synthetic assumption |",
        f"| `dusuk_verim_cezasi` | {values['low_yield_penalty']} | USD/day | NASS price + synthetic 5 kg/day gap |",
        "",
        "## Source-Derived Calculations",
        "",
        f"- NASS file: `{NASS_FILE.relative_to(PROJECT_ROOT)}`",
        f"- NASS 2025 US TOTAL monthly row count: `{values['nass_month_count']}`",
        f"- NASS average milk price: `{values['nass_avg_per_cwt']}` USD/cwt",
        f"- Conversion: `1 cwt = {CWT_TO_KG}` kg",
        f"- Milk price inserted: `{values['milk_price_usd_per_kg']}` USD/kg",
        "",
        f"- ERS file: `{ERS_STATE_FILE.relative_to(PROJECT_ROOT)}`",
        f"- ERS output per cow: `{values['ers_output_per_cow_lb_year']}` lb/year = `{values['ers_output_per_cow_cwt_year']}` cwt/year",
        f"- ERS operating cost: `{values['ers_operating_cost_per_cwt']}` USD/cwt",
        f"- Operating daily cost inserted: `{values['daily_operating_cost']}` USD/cow/day",
        f"- Reference total costs listed: `{values['daily_total_costs_listed']}` USD/cow/day",
        f"- Reference feed cost: `{values['daily_feed_cost']}` USD/cow/day",
        "",
        "## Remaining Manual Review",
        "",
        "- USDA AMS cull-cow PDFs still need a market-class parser before they can replace `cull_satis_degeri`.",
        "- Iowa State heifer PDF still needs manual table extraction before it can replace `replacement_maliyeti`.",
        "- Treatment and insemination costs are farm protocol parameters; they are intentionally synthetic in this prototype.",
    ]
    OUTPUT_MD.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_sql(values: dict[str, Decimal | str | int]) -> None:
    OUTPUT_SQL.parent.mkdir(parents=True, exist_ok=True)
    note = baseline_notes(values)
    alias_note = (
        "Legacy placeholder alias populated from baseline_us_2025_mixed_source_v1 "
        "so prototype DB checks are not left with empty economic parameters. "
        "Use baseline_us_2025_mixed_source_v1 as the canonical scenario."
    )
    common_values = {
        "sut_fiyati": values["milk_price_usd_per_kg"],
        "replacement_maliyeti": values["replacement_cost"],
        "cull_satis_degeri": values["cull_sale_value"],
        "gunluk_bakim_maliyeti": values["daily_operating_cost"],
        "tedavi_maliyeti": values["treatment_cost"],
        "basarisiz_tohumlama_maliyeti": values["failed_insemination_cost"],
        "dusuk_verim_cezasi": values["low_yield_penalty"],
    }
    lines = [
        "/*",
        "Seed reviewed reward parameters for a US 2025 prototype baseline.",
        "",
        "Source-derived values:",
        "- sut_fiyati: USDA NASS 2025 US TOTAL average milk price converted from USD/cwt to USD/kg.",
        "- gunluk_bakim_maliyeti: USDA ERS 2025 U.S. total operating cost converted to USD/cow/day.",
        "",
        "Synthetic assumptions are explicitly documented in aciklama and in",
        "reports/reward_baseline_us_2025_review.md.",
        "*/",
        "",
        "begin;",
        "",
        "drop table if exists pg_temp.stg_reward_parametre_us_2025;",
        "create temp table stg_reward_parametre_us_2025 on commit drop as",
        "select",
        f"    {sql_text(SCENARIO_CODE)}::text as senaryo_kodu,",
        f"    date '{VALID_FROM}' as gecerlilik_baslangic,",
        f"    date '{VALID_TO}' as gecerlilik_bitis,",
        f"    {common_values['sut_fiyati']}::numeric(12,4) as sut_fiyati,",
        f"    {common_values['replacement_maliyeti']}::numeric(12,2) as replacement_maliyeti,",
        f"    {common_values['cull_satis_degeri']}::numeric(12,2) as cull_satis_degeri,",
        f"    {common_values['gunluk_bakim_maliyeti']}::numeric(12,4) as gunluk_bakim_maliyeti,",
        f"    {common_values['tedavi_maliyeti']}::numeric(12,2) as tedavi_maliyeti,",
        f"    {common_values['basarisiz_tohumlama_maliyeti']}::numeric(12,2) as basarisiz_tohumlama_maliyeti,",
        f"    {common_values['dusuk_verim_cezasi']}::numeric(12,2) as dusuk_verim_cezasi,",
        f"    {sql_text(note)}::text as aciklama;",
        "",
        "update suru_karar.feat_reward_parametre target",
        "set",
        "    gecerlilik_bitis = source.gecerlilik_bitis,",
        "    sut_fiyati = source.sut_fiyati,",
        "    replacement_maliyeti = source.replacement_maliyeti,",
        "    cull_satis_degeri = source.cull_satis_degeri,",
        "    gunluk_bakim_maliyeti = source.gunluk_bakim_maliyeti,",
        "    tedavi_maliyeti = source.tedavi_maliyeti,",
        "    basarisiz_tohumlama_maliyeti = source.basarisiz_tohumlama_maliyeti,",
        "    dusuk_verim_cezasi = source.dusuk_verim_cezasi,",
        "    aciklama = source.aciklama",
        "from stg_reward_parametre_us_2025 source",
        "where target.senaryo_kodu = source.senaryo_kodu",
        "  and target.gecerlilik_baslangic = source.gecerlilik_baslangic;",
        "",
        "insert into suru_karar.feat_reward_parametre (",
        "    senaryo_kodu, gecerlilik_baslangic, gecerlilik_bitis, sut_fiyati,",
        "    replacement_maliyeti, cull_satis_degeri, gunluk_bakim_maliyeti,",
        "    tedavi_maliyeti, basarisiz_tohumlama_maliyeti, dusuk_verim_cezasi, aciklama",
        ")",
        "select",
        "    senaryo_kodu, gecerlilik_baslangic, gecerlilik_bitis, sut_fiyati,",
        "    replacement_maliyeti, cull_satis_degeri, gunluk_bakim_maliyeti,",
        "    tedavi_maliyeti, basarisiz_tohumlama_maliyeti, dusuk_verim_cezasi, aciklama",
        "from stg_reward_parametre_us_2025 source",
        "where not exists (",
        "    select 1",
        "    from suru_karar.feat_reward_parametre target",
        "    where target.senaryo_kodu = source.senaryo_kodu",
        "      and target.gecerlilik_baslangic = source.gecerlilik_baslangic",
        ");",
        "",
        "drop table if exists pg_temp.stg_reward_parametre_legacy_alias;",
        "create temp table stg_reward_parametre_legacy_alias on commit drop as",
        "select",
        f"    {sql_text(LEGACY_SCENARIO_CODE)}::text as senaryo_kodu,",
        "    date '2026-01-01' as gecerlilik_baslangic,",
        "    null::date as gecerlilik_bitis,",
        "    sut_fiyati, replacement_maliyeti, cull_satis_degeri, gunluk_bakim_maliyeti,",
        "    tedavi_maliyeti, basarisiz_tohumlama_maliyeti, dusuk_verim_cezasi,",
        f"    {sql_text(alias_note + ' ' + note)}::text as aciklama",
        "from stg_reward_parametre_us_2025;",
        "",
        "update suru_karar.feat_reward_parametre target",
        "set",
        "    gecerlilik_bitis = source.gecerlilik_bitis,",
        "    sut_fiyati = source.sut_fiyati,",
        "    replacement_maliyeti = source.replacement_maliyeti,",
        "    cull_satis_degeri = source.cull_satis_degeri,",
        "    gunluk_bakim_maliyeti = source.gunluk_bakim_maliyeti,",
        "    tedavi_maliyeti = source.tedavi_maliyeti,",
        "    basarisiz_tohumlama_maliyeti = source.basarisiz_tohumlama_maliyeti,",
        "    dusuk_verim_cezasi = source.dusuk_verim_cezasi,",
        "    aciklama = source.aciklama",
        "from stg_reward_parametre_legacy_alias source",
        "where target.senaryo_kodu = source.senaryo_kodu",
        "  and target.gecerlilik_baslangic = source.gecerlilik_baslangic;",
        "",
        "insert into suru_karar.feat_reward_parametre (",
        "    senaryo_kodu, gecerlilik_baslangic, gecerlilik_bitis, sut_fiyati,",
        "    replacement_maliyeti, cull_satis_degeri, gunluk_bakim_maliyeti,",
        "    tedavi_maliyeti, basarisiz_tohumlama_maliyeti, dusuk_verim_cezasi, aciklama",
        ")",
        "select",
        "    senaryo_kodu, gecerlilik_baslangic, gecerlilik_bitis, sut_fiyati,",
        "    replacement_maliyeti, cull_satis_degeri, gunluk_bakim_maliyeti,",
        "    tedavi_maliyeti, basarisiz_tohumlama_maliyeti, dusuk_verim_cezasi, aciklama",
        "from stg_reward_parametre_legacy_alias source",
        "where not exists (",
        "    select 1",
        "    from suru_karar.feat_reward_parametre target",
        "    where target.senaryo_kodu = source.senaryo_kodu",
        "      and target.gecerlilik_baslangic = source.gecerlilik_baslangic",
        ");",
        "",
        "commit;",
    ]
    OUTPUT_SQL.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> int:
    values = build_values()
    write_baseline_csv(values)
    write_review(values)
    write_sql(values)
    print(f"wrote {OUTPUT_CSV.relative_to(PROJECT_ROOT)}")
    print(f"wrote {OUTPUT_MD.relative_to(PROJECT_ROOT)}")
    print(f"wrote {OUTPUT_SQL.relative_to(PROJECT_ROOT)}")
    print(f"scenario={SCENARIO_CODE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
